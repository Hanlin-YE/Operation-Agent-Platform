from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.services.neo4j_client import Neo4jClient


REQUIRED_HEADERS = {
    "线索ID",
    "账号昵称",
    "账号类型",
    "应用场景",
    "信号类型",
    "意向等级",
    "跟进状态",
    "笔记标题",
    "公开证据摘要",
    "判断理由",
    "推荐跟进动作",
    "笔记链接",
    "主页链接",
    "笔记/采集时间",
}

SCENARIO_TO_PRODUCT = {
    "机器人/灵巧手": "product:fx-integrated-solution",
    "医疗精密设备": "product:fx-brushed-coreless",
    "汽车电子": "product:fx-integrated-solution",
    "无人机/航空航天": "product:fx-brushed-coreless",
    "工业自动化": "product:fx-integrated-solution",
    "全部场景": "product:fx-brushed-coreless",
}

SCENARIO_IDS = {
    "机器人/灵巧手": "scenario:robotics",
    "医疗精密设备": "scenario:medical",
    "汽车电子": "scenario:automotive-electronics",
    "无人机/航空航天": "scenario:drone-aerospace",
    "工业自动化": "scenario:industrial-automation",
    "全部场景": "scenario:xhs-general",
}


@dataclass(frozen=True)
class GraphImportPayload:
    entities: list[dict[str, Any]]
    relationships: list[dict[str, str]]


def _stable_hash(value: str) -> str:
    return hashlib.sha1(value.encode("utf-8")).hexdigest()[:16]


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _clip(value: str, limit: int = 420) -> str:
    value = " ".join(value.split())
    return value if len(value) <= limit else f"{value[: limit - 1]}…"


def _cell_link(cell: Any) -> str:
    if cell.hyperlink and cell.hyperlink.target:
        return _clean(cell.hyperlink.target)
    return _clean(cell.value)


def build_xiaohongshu_graph_payload(xlsx_path: Path) -> GraphImportPayload:
    workbook = load_workbook(xlsx_path, read_only=False, data_only=True)
    if "公开线索表" not in workbook.sheetnames:
        raise ValueError("工作簿缺少“公开线索表”工作表")

    sheet = workbook["公开线索表"]
    headers = [_clean(cell.value) for cell in sheet[1]]
    missing = sorted(REQUIRED_HEADERS - set(headers))
    if missing:
        raise ValueError(f"公开线索表缺少必要字段：{', '.join(missing)}")

    column = {header: headers.index(header) + 1 for header in headers}
    entities: dict[str, dict[str, Any]] = {
        "source:xiaohongshu-public-leads-2026-07-21": {
            "id": "source:xiaohongshu-public-leads-2026-07-21",
            "name": "小红书公开线索表 2026-07-21",
            "type": "Source",
            "platform": "小红书",
            "source_file": xlsx_path.name,
            "data_boundary": "仅公开笔记/评论快照，触达前需人工核验",
        },
        "product:fx-brushed-coreless": {
            "id": "product:fx-brushed-coreless",
            "name": "自超滑有刷空心杯电机平台",
            "type": "Product",
            "source": "已审核产品目录",
            "selection_note": "按项目选型确认",
        },
        "product:fx-integrated-solution": {
            "id": "product:fx-integrated-solution",
            "name": "电机 + 齿轮箱 + FOC 控制器整体方案",
            "type": "Product",
            "source": "已审核产品目录",
            "selection_note": "按项目选型确认",
        },
    }
    relationships: set[tuple[str, str, str]] = set()
    source_id = "source:xiaohongshu-public-leads-2026-07-21"

    for row_index in range(2, sheet.max_row + 1):
        lead_id = _clean(sheet.cell(row=row_index, column=column["线索ID"]).value)
        nickname = _clean(sheet.cell(row=row_index, column=column["账号昵称"]).value)
        if not lead_id or not nickname:
            continue

        profile_url = _cell_link(sheet.cell(row=row_index, column=column["主页链接"]))
        note_url = _cell_link(sheet.cell(row=row_index, column=column["笔记链接"]))
        scene = _clean(sheet.cell(row=row_index, column=column["应用场景"]).value) or "未分类场景"
        signal_type = _clean(sheet.cell(row=row_index, column=column["信号类型"]).value) or "未分类信号"
        intent_level = _clean(sheet.cell(row=row_index, column=column["意向等级"]).value) or "未知"
        action = _clean(sheet.cell(row=row_index, column=column["推荐跟进动作"]).value) or "人工核验后跟进"
        status = _clean(sheet.cell(row=row_index, column=column["跟进状态"]).value)
        collected_at = _clean(sheet.cell(row=row_index, column=column["笔记/采集时间"]).value)

        contact_id = f"contact:xhs:{_stable_hash(profile_url or nickname)}"
        evidence_id = f"evidence:xhs:{lead_id}"
        scenario_id = SCENARIO_IDS.get(scene, f"scenario:xhs:{_stable_hash(scene)}")
        signal_id = f"signal-type:xhs:{_stable_hash(signal_type)}"
        intent_id = f"intent:xhs:{intent_level}"
        action_id = f"action:xhs:{_stable_hash(action)}"

        entities[contact_id] = {
            "id": contact_id,
            "name": nickname,
            "type": "Contact",
            "platform": "小红书",
            "account_type": _clean(sheet.cell(row=row_index, column=column["账号类型"]).value),
            "profile_url": profile_url,
            "verification_status": "待人工核验",
        }
        entities[evidence_id] = {
            "id": evidence_id,
            "name": _clip(_clean(sheet.cell(row=row_index, column=column["笔记标题"]).value) or f"小红书线索 {row_index - 1}", 72),
            "type": "Evidence",
            "lead_id": lead_id,
            "platform": "小红书",
            "note_url": note_url,
            "profile_url": profile_url,
            "evidence_summary": _clip(_clean(sheet.cell(row=row_index, column=column["公开证据摘要"]).value)),
            "judgement_reason": _clip(_clean(sheet.cell(row=row_index, column=column["判断理由"]).value)),
            "follow_up_status": status,
            "collected_at": collected_at,
        }
        entities[scenario_id] = {
            "id": scenario_id,
            "name": scene,
            "type": "Scenario",
            "need": "由小红书公开内容识别出的应用场景",
        }
        entities[signal_id] = {"id": signal_id, "name": signal_type, "type": "SignalType"}
        entities[intent_id] = {"id": intent_id, "name": f"{intent_level}意向", "type": "IntentLevel"}
        entities[action_id] = {"id": action_id, "name": _clip(action, 60), "type": "FollowUpAction", "action_detail": action}

        relationships.update(
            {
                (source_id, evidence_id, "HAS_LEAD"),
                (contact_id, evidence_id, "PROVIDES_SIGNAL"),
                (contact_id, scenario_id, "HAS_SCENARIO"),
                (evidence_id, scenario_id, "OBSERVED_IN"),
                (evidence_id, signal_id, "HAS_SIGNAL_TYPE"),
                (evidence_id, intent_id, "HAS_INTENT"),
                (evidence_id, action_id, "SUGGESTS_ACTION"),
            }
        )
        product_id = SCENARIO_TO_PRODUCT.get(scene)
        if product_id:
            relationships.add((scenario_id, product_id, "MATCHES"))

    return GraphImportPayload(
        entities=list(entities.values()),
        relationships=[
            {"source": source, "target": target, "type": relationship_type}
            for source, target, relationship_type in sorted(relationships)
        ],
    )


async def import_xiaohongshu_leads_to_neo4j(client: Neo4jClient, xlsx_path: Path) -> dict[str, int]:
    payload = build_xiaohongshu_graph_payload(xlsx_path)
    imported_at = datetime.now(timezone.utc).isoformat()
    await client.run_cypher(
        "CREATE CONSTRAINT graph_entity_id IF NOT EXISTS FOR (entity:GraphEntity) REQUIRE entity.id IS UNIQUE",
        {},
    )
    await client.run_cypher(
        """
        UNWIND $entities AS entity
        MERGE (node:GraphEntity {id: entity.id})
        SET node += entity
        """,
        {"entities": payload.entities},
    )
    await client.run_cypher(
        """
        UNWIND $relationships AS relationship
        MATCH (source:GraphEntity {id: relationship.source}), (target:GraphEntity {id: relationship.target})
        MERGE (source)-[relation:RELATED {type: relationship.type}]->(target)
        SET relation.imported_at = $imported_at,
            relation.source = $source
        """,
        {
            "relationships": payload.relationships,
            "imported_at": imported_at,
            "source": "小红书公开线索表_2026-07-21.xlsx",
        },
    )
    return {"entities": len(payload.entities), "relationships": len(payload.relationships)}
